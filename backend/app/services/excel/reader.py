from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, List, Tuple

import openpyxl


def read_excel_bytes(
    file_bytes: bytes,
    *,
    sheet_name: str | None = None,
    header_row: int = 1,
) -> Tuple[List[str], List[Dict[str, Any]]]:
    """
    Reads an .xlsx file from bytes and returns:
      - headers: list[str]
      - rows: list[dict] where keys are headers

    Assumes first row contains headers (header_row=1).
    """
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)

    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}")
        ws = wb[sheet_name]
    else:
        ws = wb.active

    # Read header row
    headers = []
    for cell in ws[header_row]:
        val = cell.value
        headers.append(str(val).strip() if val is not None else "")

    if not any(h for h in headers):
        raise ValueError("Header row is empty or missing.")

    # Read data rows
    rows: List[Dict[str, Any]] = []
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if r is None:
            continue

        # Skip fully empty rows
        if all(v is None or (isinstance(v, str) and v.strip() == "") for v in r):
            continue

        row_dict: Dict[str, Any] = {}
        for idx, h in enumerate(headers):
            if h == "":
                continue
            row_dict[h] = r[idx] if idx < len(r) else None

        rows.append(row_dict)

    return headers, rows


def validate_required_columns(headers: List[str], required: List[str]) -> List[str]:
    header_set = {h.strip() for h in headers if h}
    missing = [c for c in required if c not in header_set]
    return missing
