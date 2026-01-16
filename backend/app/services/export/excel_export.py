from __future__ import annotations

from typing import List
from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.models.record import Record

CALC_HEADERS = ["New Daily budget", "New target CPA"]


def records_to_excel(records: List[Record]) -> BytesIO:
    if not records:
        raise ValueError("No records to export")

    # İlk header'ı DB’den bul
    source_headers = None
    for r in records:
        if r.source_headers:
            source_headers = list(r.source_headers)
            break

    # Fallback: header yoksa raw_row keys
    if not source_headers:
        for r in records:
            if r.raw_row:
                source_headers = list(r.raw_row.keys())
                break

    if not source_headers:
        raise ValueError("No source headers found. Upload must store headers/raw_row.")

    # Final header: orijinal + 2 yeni kolon (eğer orijinalde yoksa ekle)
    headers = list(source_headers)
    for h in CALC_HEADERS:
        if h not in headers:
            headers.append(h)

    wb = Workbook()
    ws = wb.active
    ws.title = "Calculated Data"

    ws.append(headers)

    for r in records:
        row = dict(r.raw_row or {})

        # raw_row'da yoksa DB field'lardan doldur
        row.setdefault("Campaign name", r.campaign_name)
        row.setdefault("Labels on Campaign", r.label)

        row.setdefault("Camp. budget", r.camp_budget)
        row.setdefault("Camp. cost", r.camp_cost)
        row.setdefault("Camp. 3D cost", r.camp_3d_cost)
        row.setdefault("Camp. conv.", r.camp_conv)
        row.setdefault("Camp. CPA", r.camp_cpa)
        row.setdefault("Camp. tCPA", r.camp_tcpa)

        row.setdefault("MTD Cluster Spend", r.mtd_cluster_spend)

        row.setdefault("Label budget", r.label_budget)
        row.setdefault("Label cost", r.label_cost)
        row.setdefault("Label 3D cost", r.label_3d_cost)
        row.setdefault("Label conv.", r.label_conv)
        row.setdefault("Label remaining budget", r.label_remaining_budget)
        row.setdefault("Label KPI value", r.label_kpi_value)
        row.setdefault("Label CPA", r.label_cpa)

        # Hesaplanan kolonlar
        row["New Daily budget"] = r.new_daily_budget
        row["New target CPA"] = r.new_target_cpa

        ws.append([row.get(h) for h in headers])

    for i, col in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(len(str(col)) + 2, 18)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output